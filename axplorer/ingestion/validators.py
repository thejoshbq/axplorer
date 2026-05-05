"""Input validation for signal and event files.

Provides pre-flight checks for neural signal (.npy) and behavioral event
(.xlsx / .mat) files before they are handed to the Pynapse Sample constructor.
Each validator returns a ``ValidationResult`` containing accumulated errors
and warnings so callers can decide how to proceed.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import List

import numpy as np


# ---------------------------------------------------------------------------
# Result container
# ---------------------------------------------------------------------------

@dataclass
class ValidationResult:
    """Result of a file validation check.

    Attributes:
        valid: ``True`` if no blocking errors were found.
        errors: Descriptions of problems that prevent loading.
        warnings: Non-blocking issues the user should know about.
    """

    valid: bool
    errors: List[str] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Signal file (.npy)
# ---------------------------------------------------------------------------

def validate_signal_file(path: str | Path) -> ValidationResult:
    """Validate a ``.npy`` neural signal file.

    Checks performed (in order):
        1. File exists on disk.
        2. Extension is ``.npy``.
        3. File is loadable by ``numpy.load``.
        4. Array is two-dimensional (neurons x frames).
        5. ``shape[0]`` (neurons) < ``shape[1]`` (frames).
        6. Array contains no NaN or Inf values.

    Args:
        path: Filesystem path to the signal file.

    Returns:
        A ``ValidationResult`` with accumulated errors and warnings.
    """
    errors: List[str] = []
    warnings: List[str] = []
    path = Path(path)

    # 1. Existence
    if not path.exists():
        return ValidationResult(valid=False, errors=[f"Signal file not found: {path}"])

    # 2. Extension
    if path.suffix.lower() != ".npy":
        errors.append(f"Expected .npy extension, got '{path.suffix}'.")

    # 3. Loadable
    try:
        data = np.load(path, allow_pickle=False)
    except Exception as exc:
        errors.append(f"Failed to load .npy file: {exc}")
        return ValidationResult(valid=False, errors=errors, warnings=warnings)

    # 4. Squeeze singleton dimensions (e.g. (1, neurons, frames) → (neurons, frames))
    if data.ndim != 2:
        squeezed = data.squeeze()
        if squeezed.ndim == 2:
            warnings.append(
                f"Signal array had shape {data.shape}; squeezed singleton "
                f"dimension(s) to {squeezed.shape}."
            )
            data = squeezed
        else:
            errors.append(
                f"Signal array must be 2-D (neurons x frames), got {data.ndim}-D "
                f"with shape {data.shape}."
            )
            return ValidationResult(valid=False, errors=errors, warnings=warnings)

    # 5. Orientation: neurons < frames
    n_neurons, n_frames = data.shape
    if n_neurons >= n_frames:
        errors.append(
            f"Expected neurons (axis 0) < frames (axis 1), got shape "
            f"({n_neurons}, {n_frames}). Is the array transposed?"
        )

    # 6. NaN / Inf
    if np.any(np.isnan(data)):
        errors.append("Signal array contains NaN values.")
    if np.any(np.isinf(data)):
        errors.append("Signal array contains Inf values.")

    return ValidationResult(
        valid=len(errors) == 0,
        errors=errors,
        warnings=warnings,
    )


# ---------------------------------------------------------------------------
# Event file (.csv / .xlsx / .mat)
# ---------------------------------------------------------------------------

_REQUIRED_EVENT_COLUMNS = {"device", "event", "start_timestamp", "end_timestamp"}


def validate_event_file(path: str | Path) -> ValidationResult:
    """Validate a behavioral event file (``.csv``, ``.xlsx``, or ``.mat``).

    For ``.csv`` files (REACHER / Labrynth export):
        * Checks that the required columns (device, event, start_timestamp,
          end_timestamp) are present in the header.

    For ``.xlsx`` files:
        * Checks that a sheet named ``"Behavior Data"`` exists.
        * Checks that the required columns are present.

    For ``.mat`` files:
        * Checks that the ``'eventlog'`` key exists in the loaded struct.

    Args:
        path: Filesystem path to the event file.

    Returns:
        A ``ValidationResult`` with accumulated errors and warnings.
    """
    errors: List[str] = []
    warnings: List[str] = []
    path = Path(path)

    # Existence
    if not path.exists():
        return ValidationResult(valid=False, errors=[f"Event file not found: {path}"])

    suffix = path.suffix.lower()

    if suffix == ".csv":
        errors, warnings = _validate_csv_event_file(path)
    elif suffix == ".xlsx":
        errors, warnings = _validate_xlsx_event_file(path)
    elif suffix == ".mat":
        errors, warnings = _validate_mat_event_file(path)
    else:
        errors.append(
            f"Unsupported event file extension '{suffix}'. "
            "Expected .csv, .xlsx, or .mat."
        )

    return ValidationResult(
        valid=len(errors) == 0,
        errors=errors,
        warnings=warnings,
    )


def _validate_csv_event_file(path: Path) -> tuple[List[str], List[str]]:
    """Validate a REACHER-exported ``behavior_events.csv`` file.

    Only the header row is parsed so that very large CSVs remain cheap to
    check. The required column set is the same as for ``.xlsx``.

    Args:
        path: Path to the .csv file.

    Returns:
        Tuple of (errors, warnings).
    """
    import csv

    errors: List[str] = []
    warnings: List[str] = []

    try:
        with path.open("r", newline="", encoding="utf-8") as f:
            reader = csv.reader(f)
            try:
                header = next(reader)
            except StopIteration:
                return [f"Event CSV is empty: {path}"], warnings
    except Exception as exc:
        return [f"Failed to read .csv file: {exc}"], warnings

    found = {str(col).strip().lower() for col in header if col}
    missing = _REQUIRED_EVENT_COLUMNS - found
    if missing:
        errors.append(
            f"CSV is missing required columns: {sorted(missing)}. "
            f"Found: {sorted(found)}"
        )

    return errors, warnings


def _validate_xlsx_event_file(path: Path) -> tuple[List[str], List[str]]:
    """Validate an Excel-based event file.

    Args:
        path: Path to the .xlsx file.

    Returns:
        Tuple of (errors, warnings).
    """
    import openpyxl

    errors: List[str] = []
    warnings: List[str] = []

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        return [f"Failed to open .xlsx file: {exc}"], warnings

    if "Behavior Data" not in wb.sheetnames:
        errors.append(
            f"Missing required sheet 'Behavior Data'. "
            f"Found sheets: {wb.sheetnames}"
        )
        wb.close()
        return errors, warnings

    ws = wb["Behavior Data"]
    # Read header row (row 1)
    header = [
        str(cell.value).strip().lower() if cell.value is not None else ""
        for cell in next(ws.iter_rows(min_row=1, max_row=1))
    ]
    wb.close()

    found = set(header)
    missing = _REQUIRED_EVENT_COLUMNS - found
    if missing:
        errors.append(
            f"Sheet 'Behavior Data' is missing required columns: "
            f"{sorted(missing)}. Found: {sorted(found)}"
        )

    return errors, warnings


def _validate_mat_event_file(path: Path) -> tuple[List[str], List[str]]:
    """Validate a MATLAB .mat event file.

    Args:
        path: Path to the .mat file.

    Returns:
        Tuple of (errors, warnings).
    """
    import scipy.io as sio

    errors: List[str] = []
    warnings: List[str] = []

    try:
        mat = sio.loadmat(str(path))
    except Exception as exc:
        return [f"Failed to load .mat file: {exc}"], warnings

    if "eventlog" not in mat:
        available = [k for k in mat if not k.startswith("__")]
        errors.append(
            f"Missing required key 'eventlog' in .mat file. "
            f"Available keys: {sorted(available)}"
        )

    return errors, warnings


# ---------------------------------------------------------------------------
# Task-type detection
# ---------------------------------------------------------------------------

def detect_task_type(path: str | Path) -> str:
    """Auto-detect task type from an event file.

    Detection rules:
        * ``.csv`` files originate from the REACHER/Labrynth export pipeline.
        * ``.xlsx`` files are assumed to originate from the Reacher paradigm.
        * ``.mat`` files containing event codes 50 or 51 in the ``eventlog``
          are classified as ``'legacy_eth'`` (ethanol self-administration).
        * All other ``.mat`` files default to ``'legacy_her'`` (heroin
          self-administration).

    Args:
        path: Filesystem path to the event file.

    Returns:
        One of ``'reacher'``, ``'legacy_eth'``, or ``'legacy_her'``.

    Raises:
        ValueError: If the file extension is unsupported or the .mat file
            cannot be read.
    """
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix in (".csv", ".xlsx"):
        return "reacher"

    if suffix == ".mat":
        import scipy.io as sio

        try:
            mat = sio.loadmat(str(path))
        except Exception as exc:
            raise ValueError(f"Cannot read .mat file for task detection: {exc}") from exc

        eventlog = mat.get("eventlog")
        if eventlog is None:
            raise ValueError("Cannot detect task type: 'eventlog' key missing from .mat file.")

        event_codes = np.squeeze(eventlog)
        if event_codes.ndim == 2:
            event_codes = event_codes[:, 0]

        eth_codes = {50, 51}
        if eth_codes & set(event_codes.astype(int)):
            return "legacy_eth"
        return "legacy_her"

    raise ValueError(
        f"Cannot detect task type for extension '{suffix}'. "
        "Expected .csv, .xlsx, or .mat."
    )


# ---------------------------------------------------------------------------
# Frame timestamps file (.csv)
# ---------------------------------------------------------------------------

_REQUIRED_FRAME_TS_COLUMNS = {"timestamp_ms"}


def validate_frame_timestamps_file(path: str | Path) -> ValidationResult:
    """Validate a ``frame_timestamps.csv`` file produced by REACHER.

    Checks performed:
        1. File exists.
        2. Extension is ``.csv``.
        3. Header contains ``timestamp_ms`` (required by pynapse's
           ``Sample(frame_timestamps=...)`` loader).
        4. A ``frame_index`` column is recommended (warned when missing; pynapse
           tolerates its absence by assigning implicit ordering).

    Args:
        path: Filesystem path to the frame timestamps file.

    Returns:
        A ``ValidationResult`` with accumulated errors and warnings.
    """
    import csv

    errors: List[str] = []
    warnings: List[str] = []
    path = Path(path)

    if not path.exists():
        return ValidationResult(valid=False, errors=[f"Frame timestamps file not found: {path}"])

    if path.suffix.lower() != ".csv":
        errors.append(f"Expected .csv extension, got '{path.suffix}'.")
        return ValidationResult(valid=False, errors=errors, warnings=warnings)

    try:
        with path.open("r", newline="", encoding="utf-8") as f:
            reader = csv.reader(f)
            try:
                header = next(reader)
            except StopIteration:
                return ValidationResult(
                    valid=False,
                    errors=[f"Frame timestamps CSV is empty: {path}"],
                )
    except Exception as exc:
        return ValidationResult(
            valid=False,
            errors=[f"Failed to read frame timestamps file: {exc}"],
        )

    found = {str(col).strip().lower() for col in header if col}
    missing = _REQUIRED_FRAME_TS_COLUMNS - found
    if missing:
        errors.append(
            f"frame_timestamps.csv is missing required columns: {sorted(missing)}. "
            f"Found: {sorted(found)}"
        )

    if "frame_index" not in found:
        warnings.append(
            "frame_timestamps.csv has no 'frame_index' column; pynapse will "
            "assign implicit ordering based on row order."
        )

    return ValidationResult(
        valid=len(errors) == 0,
        errors=errors,
        warnings=warnings,
    )


# ---------------------------------------------------------------------------
# Frame alignment check
# ---------------------------------------------------------------------------

def validate_alignment(
    n_signal_frames: int,
    n_timestamp_frames: int,
    frame_averaging: int,
) -> ValidationResult:
    """Check that signal frame count aligns with timestamp frame count.

    After dividing the raw timestamp frame count by ``frame_averaging``, the
    result should match ``n_signal_frames``. Small discrepancies (1--2 frames)
    are typical due to acquisition start/stop edge effects and produce a
    warning. Larger discrepancies indicate a likely configuration mismatch.

    Args:
        n_signal_frames: Number of frames in the signal array (axis 1).
        n_timestamp_frames: Number of raw frame-trigger timestamps extracted
            from the event log.
        frame_averaging: Number of raw frames binned into each signal frame.

    Returns:
        A ``ValidationResult`` with any alignment warnings or errors.
    """
    errors: List[str] = []
    warnings: List[str] = []

    expected = n_timestamp_frames // frame_averaging
    diff = abs(n_signal_frames - expected)

    if diff == 0:
        pass  # perfect alignment
    elif diff <= 2:
        warnings.append(
            f"Minor frame mismatch: signal has {n_signal_frames} frames, "
            f"expected {expected} from {n_timestamp_frames} timestamps "
            f"/ {frame_averaging} averaging (off by {diff})."
        )
    elif diff <= 5:
        warnings.append(
            f"Notable frame mismatch: signal has {n_signal_frames} frames, "
            f"expected {expected} from {n_timestamp_frames} timestamps "
            f"/ {frame_averaging} averaging (off by {diff}). "
            "Check acquisition logs."
        )
    else:
        errors.append(
            f"Severe frame mismatch: signal has {n_signal_frames} frames, "
            f"expected {expected} from {n_timestamp_frames} timestamps "
            f"/ {frame_averaging} averaging (off by {diff}). "
            "Verify signal file, event file, and frame_averaging are correct."
        )

    return ValidationResult(
        valid=len(errors) == 0,
        errors=errors,
        warnings=warnings,
    )
