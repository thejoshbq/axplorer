"""Shared test fixtures — mock signals, mock event files, loaded sessions."""

from __future__ import annotations

import tempfile
from pathlib import Path

import numpy as np
import openpyxl
import pytest


# ──────────────────────────────────────────────────────────────────────────
# Constants
# ──────────────────────────────────────────────────────────────────────────

N_NEURONS = 10
N_FRAMES = 5000
FPS = 30.0
FRAME_AVG = 4


# ──────────────────────────────────────────────────────────────────────────
# Signal fixtures
# ──────────────────────────────────────────────────────────────────────────

@pytest.fixture()
def mock_signal_path(tmp_path: Path) -> Path:
    """Create a synthetic .npy signal file (N_NEURONS x N_FRAMES)."""
    rng = np.random.default_rng(42)
    signals = rng.normal(loc=100, scale=10, size=(N_NEURONS, N_FRAMES)).astype(np.float32)
    path = tmp_path / "mock_signals.npy"
    np.save(path, signals)
    return path


@pytest.fixture()
def bad_signal_1d(tmp_path: Path) -> Path:
    """Create a 1-D .npy file (invalid for signals)."""
    path = tmp_path / "bad_1d.npy"
    np.save(path, np.zeros(100, dtype=np.float32))
    return path


@pytest.fixture()
def signal_with_nan(tmp_path: Path) -> Path:
    """Create a signal file containing NaN values."""
    data = np.ones((5, 500), dtype=np.float32)
    data[2, 100] = np.nan
    path = tmp_path / "nan_signal.npy"
    np.save(path, data)
    return path


# ──────────────────────────────────────────────────────────────────────────
# Event fixtures
# ──────────────────────────────────────────────────────────────────────────

def _make_reacher_xlsx(path: Path, n_events: int = 50) -> Path:
    """Build a minimal REACHER-format .xlsx with Behavior Data + Frame Timestamps."""
    wb = openpyxl.Workbook()

    # -- Session Summary (not validated, but present for completeness) ------
    ws_summary = wb.active
    ws_summary.title = "Session Summary"

    # -- Behavior Data ------------------------------------------------------
    # Column A = row index (read by Pynapse with index_col=0),
    # then device, event, start_timestamp, end_timestamp.
    ws_beh = wb.create_sheet("Behavior Data")
    ws_beh.append(["index", "device", "event", "start_timestamp", "end_timestamp"])

    rng = np.random.default_rng(99)
    devices_events = [
        ("RH_LEVER", "ACTIVE_PRESS"),
        ("LH_LEVER", "INACTIVE_PRESS"),
        ("RH_LEVER", "TIMEOUT_PRESS"),
        ("PUMP", "INFUSION"),
        ("CUE", "TONE"),
    ]

    # Timestamps are in ms (matching Pynapse's interframe_interval = 1000/fps).
    # Signal is N_FRAMES=5000 averaged frames at effective_fps=7.5 → ~666s.
    # Place events well within signal bounds: start at 30s, space by 5-10s.
    ifi_ms = 1000.0 / FPS  # ~33.333 ms
    total_duration_ms = N_FRAMES * FRAME_AVG * ifi_ms  # full raw session in ms
    ts = 30_000.0  # 30s in ms
    for i in range(n_events):
        dev, evt = devices_events[i % len(devices_events)]
        ts += float(rng.integers(5000, 10000))  # 5-10s spacing in ms
        if ts > total_duration_ms - 30_000:  # stop 30s before end
            break
        ws_beh.append([i, dev, evt, ts, ts + 100])

    # -- Frame Timestamps ---------------------------------------------------
    # Column A = row index (read by Pynapse with index_col=0),
    # then timestamp in ms (matching interframe_interval).
    ws_frames = wb.create_sheet("Frame Timestamps")
    ws_frames.append(["index", "timestamp"])
    frame_ts = 0.0
    for j in range(N_FRAMES * FRAME_AVG):
        ws_frames.append([j, frame_ts])
        frame_ts += ifi_ms

    wb.save(path)
    return path


@pytest.fixture()
def mock_event_path(tmp_path: Path) -> Path:
    """Create a minimal REACHER-format .xlsx event file."""
    return _make_reacher_xlsx(tmp_path / "mock_events.xlsx")


@pytest.fixture()
def bad_event_xlsx(tmp_path: Path) -> Path:
    """Create an .xlsx missing the 'Behavior Data' sheet."""
    wb = openpyxl.Workbook()
    wb.active.title = "Wrong Sheet"
    path = tmp_path / "bad_events.xlsx"
    wb.save(path)
    return path


def _make_reacher_csv_pair(dir_: Path, n_events: int = 50) -> tuple[Path, Path]:
    """Write REACHER-style behavior_events.csv + frame_timestamps.csv."""
    import csv

    dir_.mkdir(parents=True, exist_ok=True)
    ifi_ms = 1000.0 / FPS
    total_duration_ms = N_FRAMES * FRAME_AVG * ifi_ms

    rng = np.random.default_rng(99)
    devices_events = [
        ("RH_LEVER", "ACTIVE_PRESS"),
        ("LH_LEVER", "INACTIVE_PRESS"),
        ("RH_LEVER", "TIMEOUT_PRESS"),
        ("PUMP", "INFUSION"),
        ("CUE", "TONE"),
    ]

    behavior_path = dir_ / "behavior_events.csv"
    with behavior_path.open("w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["device", "event", "start_timestamp", "end_timestamp"])
        ts = 30_000.0
        for i in range(n_events):
            dev, evt = devices_events[i % len(devices_events)]
            ts += float(rng.integers(5000, 10000))
            if ts > total_duration_ms - 30_000:
                break
            writer.writerow([dev, evt, ts, ts + 100])

    ft_path = dir_ / "frame_timestamps.csv"
    with ft_path.open("w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["frame_index", "timestamp_ms"])
        frame_ts = 0.0
        for j in range(N_FRAMES * FRAME_AVG):
            writer.writerow([j, frame_ts])
            frame_ts += ifi_ms

    return behavior_path, ft_path


@pytest.fixture()
def reacher_csv_event_file(tmp_path: Path) -> Path:
    """Return a REACHER-format behavior_events.csv (with frame_timestamps.csv sibling)."""
    behavior_path, _ = _make_reacher_csv_pair(tmp_path / "reacher_csv")
    return behavior_path


@pytest.fixture()
def reacher_frame_timestamps_file(tmp_path: Path) -> Path:
    """Return a REACHER-format frame_timestamps.csv in its own dir."""
    _, ft_path = _make_reacher_csv_pair(tmp_path / "reacher_ft")
    return ft_path


# ──────────────────────────────────────────────────────────────────────────
# Loaded session fixture
# ──────────────────────────────────────────────────────────────────────────

@pytest.fixture()
def loaded_session(mock_signal_path: Path, mock_event_path: Path):
    """Load a mock session and return (sample, metadata, wrapper)."""
    from axplorer.alignment.session import SessionWrapper
    from axplorer.ingestion.loader import load_session_files

    sample, metadata = load_session_files(
        signal_path=mock_signal_path,
        event_path=mock_event_path,
        fps=FPS,
        frame_averaging=FRAME_AVG,
        task_type="reacher",
    )
    wrapper = SessionWrapper(sample)
    return sample, metadata, wrapper


# ──────────────────────────────────────────────────────────────────────────
# Directory hierarchy fixture
# ──────────────────────────────────────────────────────────────────────────

@pytest.fixture()
def mock_data_hierarchy(tmp_path: Path) -> Path:
    """Build a temporary 3-level data directory with synthetic files.

    Structure::

        tmp_path/
          0 EarlyAcq/
            PrL-NAc-G6-1F/
              FOV1_tracked/
                signals_extractedsignals_raw.npy
                events.mat
            PrL-NAc-G6-2M/
              FOV1/
                signals_extractedsignals_raw.npy
                events.mat

    Returns:
        The root path of the hierarchy.
    """
    for animal, fov in [
        ("PrL-NAc-G6-1F", "FOV1_tracked"),
        ("PrL-NAc-G6-2M", "FOV1"),
    ]:
        fov_dir = tmp_path / "0 EarlyAcq" / animal / fov
        fov_dir.mkdir(parents=True)
        np.save(
            fov_dir / "signals_extractedsignals_raw.npy",
            np.zeros((5, 100), dtype=np.float32),
        )
        (fov_dir / "events.mat").write_bytes(b"\x00" * 128)

    return tmp_path
